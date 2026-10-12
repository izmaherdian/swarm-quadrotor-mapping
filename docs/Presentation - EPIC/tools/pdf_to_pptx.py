#!/usr/bin/env python3
"""Bungkus main.pdf jadi .pptx: tiap halaman jadi gambar penuh, video MP4 asli
ditempel di atas poster-nya.

Cadangan untuk laptop panitia yang hanya punya PowerPoint. Tata letak TIDAK
dibuat di sini -- semuanya sudah ditentukan main.tex. Skrip ini hanya:

  1. membaca anotasi /Launch dari main.pdf untuk tahu klip mana ada di halaman
     mana, dan di kotak mana persisnya (geometri hasil kompilasi, bukan tebakan),
  2. merender tiap halaman jadi PNG 300 dpi,
  3. menyusunnya jadi slide dan menaruh objek video di kotak yang sama.

Pakai:  python3 tools/pdf_to_pptx.py main.pdf main.pptx
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

DPI = 300
SW, SH = Inches(13.333), Inches(7.5)


def _mutool(pdf: Path, *args: str) -> str:
    return subprocess.run(["mutool", "show", str(pdf), *args],
                          capture_output=True, text=True, check=True).stdout


def video_boxes(pdf: Path) -> dict[int, list[tuple[str, tuple[float, ...]]]]:
    """{halaman: [(nama klip, (x0,y0,x1,y1) dalam pt), ...]}

    Dibaca dari anotasi /Launch yang ditulis \\href{run:...} di mode link.
    """
    pages = {int(n): int(o) for n, o in
             re.findall(r"page (\d+) = (\d+) 0 R", _mutool(pdf, "pages"))}
    out: dict[int, list] = {}
    for page, obj in pages.items():
        body = _mutool(pdf, str(obj))
        m = re.search(r"/Annots\s*\[([^\]]*)\]", body)
        if not m:
            continue
        for ref in re.findall(r"(\d+) 0 R", m.group(1)):
            a = _mutool(pdf, ref)
            f = re.search(r"/F\s*\(([^)]*\.mp4)\)", a)
            r = re.search(r"/Rect\s*\[\s*([\d.+-]+)\s+([\d.+-]+)\s+"
                          r"([\d.+-]+)\s+([\d.+-]+)", a)
            if f and r:
                out.setdefault(page, []).append(
                    (Path(f.group(1)).stem, tuple(float(x) for x in r.groups())))
    return out


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    pdf, out = Path(sys.argv[1]), Path(sys.argv[2])
    deck = pdf.resolve().parent

    size = _mutool(pdf, "pages").count("page ")
    box = re.search(r"/MediaBox\s*\[\s*0\s+0\s+([\d.]+)\s+([\d.]+)",
                    _mutool(pdf, str(int(re.search(r"page 1 = (\d+)",
                            _mutool(pdf, "pages")).group(1)))))
    pw, ph = (float(box.group(1)), float(box.group(2))) if box else (453.543, 255.118)
    scale = SW / pw                      # pt PDF -> EMU slide

    boxes = video_boxes(pdf)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["pdftoppm", "-r", str(DPI), "-png", str(pdf),
                        str(Path(td) / "p")], check=True)
        pngs = sorted(Path(td).glob("p-*.png"))
        if len(pngs) != size:
            raise RuntimeError(f"{len(pngs)} PNG untuk {size} halaman")

        for i, png in enumerate(pngs, start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(png), 0, 0, width=SW, height=SH)
            for name, (x0, y0, x1, y1) in boxes.get(i, []):
                clip = deck / "assets" / "clips" / f"{name}.mp4"
                poster = deck / "assets" / "posters" / f"{name}.jpg"
                if not clip.is_file():
                    raise FileNotFoundError(f"klip tidak ada: {clip}")
                # PDF berasal dari kiri-bawah, PowerPoint dari kiri-atas.
                slide.shapes.add_movie(
                    str(clip), Emu(int(x0 * scale)), Emu(int((ph - y1) * scale)),
                    Emu(int((x1 - x0) * scale)), Emu(int((y1 - y0) * scale)),
                    poster_frame_image=str(poster) if poster.is_file() else None,
                    mime_type="video/mp4")

    prs.save(out)
    n_vid = sum(len(v) for v in boxes.values())
    print(f"OK  {out.name}: {size} slide, {n_vid} video tertanam, "
          f"{out.stat().st_size / 1e6:.0f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
